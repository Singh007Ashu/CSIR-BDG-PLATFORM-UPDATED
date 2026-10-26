#!/usr/bin/env python3
"""Rank CSIR technologies against an extracted industry requirement."""
import json
import base64
import io
import math
import re
import sys

STOP = set("""the a an and or of to for with in on by from is are be as at this that these those
technology technologies industry need needs required requirement requirements solution system
project csir lab laboratory product process service using based""".split())

def tokens(value):
    words = re.findall(r"[a-z0-9][a-z0-9+./-]{1,}", str(value or "").lower())
    return [word.strip("./-") for word in words if len(word.strip("./-")) > 2 and word not in STOP]

def readiness(row):
    trl = float(row.get("trl") or 0)
    stage = str(row.get("commercialisationStage") or "").lower()
    ip = str(row.get("ipStatus") or "").lower()
    score = min(25.0, trl * 2.5)
    if any(term in stage for term in ("commercial", "transfer", "market", "ready", "production", "licensed")):
        score += 12
    if any(term in ip for term in ("granted", "filed", "protected", "patent")):
        score += 5
    label = "Ready to sell / transfer" if trl >= 7 or any(term in stage for term in ("commercial", "transfer", "market", "production")) else "Needs technical/commercial review"
    return score, label

def main():
    payload = json.load(sys.stdin)
    if payload.get("mode") == "extract":
        import openpyxl
        raw = base64.b64decode(payload.get("base64") or "")
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheets = []
        text_blocks = []
        for worksheet in workbook.worksheets:
            rows = []
            for values in worksheet.iter_rows(values_only=True):
                cleaned = ["" if value is None else str(value) for value in values[:80]]
                if any(cleaned): rows.append(cleaned)
                if len(rows) >= 80: break
            header = next((row for row in rows if sum(bool(value) for value in row) >= 2), [])
            sheets.append({"name": worksheet.title, "rows": worksheet.max_row, "columns": worksheet.max_column, "headers": header, "sampleRows": rows[:30]})
            text_blocks.append(worksheet.title + "\n" + "\n".join(" | ".join(row) for row in rows[:40]))
        print(json.dumps({"kind": "SPREADSHEET", "sheets": sheets, "extractedText": "\n\n".join(text_blocks)[:300000]}, ensure_ascii=False))
        return
    need = " ".join([str(payload.get("title") or ""), str(payload.get("description") or ""), str(payload.get("text") or "")])
    need_tokens = tokens(need)
    need_set = set(need_tokens)
    frequencies = {token: need_tokens.count(token) for token in need_set}
    matches = []
    for row in payload.get("technologies") or []:
        corpus = " ".join(str(row.get(key) or "") for key in ("name", "briefDescription", "advantages", "applications", "keyFeatures", "theme", "sector", "technologyType"))
        tech_tokens = set(tokens(corpus))
        common = need_set.intersection(tech_tokens)
        lexical = sum(1.0 + math.log1p(frequencies.get(token, 1)) for token in common)
        coverage = lexical / max(1.0, sum(1.0 + math.log1p(value) for value in frequencies.values()))
        ready_score, ready_label = readiness(row)
        total = min(100.0, coverage * 78.0 + ready_score)
        if common or total >= 18:
            reasons = []
            if common:
                reasons.append("Requirement terms matched: " + ", ".join(sorted(common, key=lambda x: frequencies.get(x, 0), reverse=True)[:8]))
            reasons.append(ready_label)
            matches.append({
                "technologyId": row.get("id"),
                "name": row.get("name"),
                "labCode": (row.get("lab") or {}).get("code"),
                "labName": (row.get("lab") or {}).get("name"),
                "score": round(total, 1),
                "matchedTerms": sorted(common)[:15],
                "readiness": ready_label,
                "trl": row.get("trl"),
                "commercialisationStage": row.get("commercialisationStage"),
                "reason": " · ".join(reasons),
            })
    matches.sort(key=lambda item: (-item["score"], -(float(item.get("trl") or 0)), str(item.get("name") or "")))
    print(json.dumps({"engine": "PYTHON_MATCHER_V1", "needKeywords": sorted(need_set, key=lambda x: frequencies.get(x, 0), reverse=True)[:30], "matches": matches[:15]}, ensure_ascii=False))

if __name__ == "__main__":
    main()
